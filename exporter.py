"""Convert recognised text into real computer files (TXT, DOCX, PDF)."""
import io
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import ListFlowable, ListItem, SimpleDocTemplate, Paragraph, Spacer


def to_txt(text: str) -> bytes:
    return text.encode("utf-8")


def to_docx(text: str, title="Converted Handwritten Document") -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    for para in text.split("\n\n") if "\n\n" in text else [text]:
        p = doc.add_paragraph(para.replace("\n", " ") if "\n\n" in text else para)
        for r in p.runs:
            r.font.name, r.font.size = "Calibri", Pt(12)
    buf = io.BytesIO(); doc.save(buf)
    return buf.getvalue()


def to_pdf(text: str, title="Converted Handwritten Document") -> bytes:
    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    for line in text.split("\n"):
        story.append(Paragraph(escape(line) or "&nbsp;", styles["BodyText"]))
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()


def to_professional_txt(title: str, items: str) -> bytes:
    lines = [line.strip() for line in items.splitlines() if line.strip()]
    content = ([title.strip(), ""] if title.strip() else [])
    content.extend(f"• {line}" for line in lines)
    return "\n".join(content).encode("utf-8")


def to_professional_docx(title: str, items: str) -> bytes:
    doc = Document()
    if title.strip():
        doc.add_heading(title.strip(), level=1)
    for line in items.splitlines():
        if line.strip():
            doc.add_paragraph(line.strip(), style="List Bullet")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def to_professional_pdf(title: str, items: str) -> bytes:
    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    story = []
    if title.strip():
        story.extend([Paragraph(escape(title.strip()), styles["Title"]), Spacer(1, 12)])
    list_items = [
        ListItem(Paragraph(escape(line.strip()), styles["BodyText"]))
        for line in items.splitlines()
        if line.strip()
    ]
    if list_items:
        story.append(ListFlowable(list_items, bulletType="bullet", leftIndent=18))
    SimpleDocTemplate(buf, pagesize=A4).build(story)
    return buf.getvalue()
